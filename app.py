from __future__ import annotations

import io
import json
import os
import subprocess
import sys
import time
from datetime import date, datetime, time as dt_time
from pathlib import Path
from urllib.parse import urlsplit

import streamlit as st

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "output"
CONFIG = ROOT / "config.yaml"
OUTPUT.mkdir(parents=True, exist_ok=True)

st.set_page_config(
    page_title="Traffic Analytics | Intelligence Dashboard",
    page_icon="🚦",
    layout="wide",
    initial_sidebar_state="expanded",
)

# --------------------------- Design system ---------------------------

st.markdown("""
<style>
:root {
    --ta-bg: #080E19;
    --ta-surface: #101A2B;
    --ta-surface-2: #17243A;
    --ta-border: #2A3B54;
    --ta-text: #EDF4FF;
    --ta-muted: #A9B8CC;
    --ta-accent: #51D7F5;
    --ta-accent-text: #07151C;
    --ta-success: #72E8BB;
    --ta-shadow: 0 10px 30px rgba(0,0,0,.16);
    --ta-radius: 18px;
}

/* ---------- Global layout ---------- */
html, body, [class*="css"] {
    font-family: Inter, -apple-system, BlinkMacSystemFont,
                 "Segoe UI", sans-serif;
}

.stApp,
[data-testid="stAppViewContainer"],
[data-testid="stMain"],
[data-testid="stMainBlockContainer"] {
    background: var(--ta-bg) !important;
    color: var(--ta-text) !important;
}

[data-testid="stMainBlockContainer"] {
    max-width: 1500px;
    padding-top: 1.6rem !important;
    padding-bottom: 3rem !important;
}

[data-testid="stHeader"] {
    background: var(--ta-bg) !important;
    border-bottom: 1px solid var(--ta-border);
    z-index: 1000;
}

[data-testid="stToolbar"] {
    background: transparent !important;
}

/* ---------- Typography ---------- */
h1, h2, h3, h4, p, label, li,
[data-testid="stMarkdownContainer"] {
    color: var(--ta-text);
}

h1 {
    font-size: clamp(1.8rem, 3vw, 2.35rem) !important;
    font-weight: 800 !important;
    letter-spacing: -.045em;
    line-height: 1.2 !important;
}

h2 {
    font-size: 1.45rem !important;
    font-weight: 750 !important;
    letter-spacing: -.025em;
}

h3 {
    font-size: 1.08rem !important;
    font-weight: 700 !important;
}

p {
    line-height: 1.65;
}

small,
.muted,
[data-testid="stCaptionContainer"] {
    color: var(--ta-muted) !important;
}

/* ---------- Sidebar ---------- */
[data-testid="stSidebar"],
[data-testid="stSidebar"] > div:first-child {
    background: var(--ta-surface) !important;
    border-right: 1px solid var(--ta-border);
}

[data-testid="stSidebar"] * {
    color: var(--ta-text);
}

[data-testid="stSidebar"] hr {
    border-color: var(--ta-border);
}

[data-testid="stRadio"] label {
    border-radius: 10px;
}

[data-testid="stRadio"] label:hover {
    background: var(--ta-surface-2);
}

/* ---------- Overview hero ---------- */
.hero {
    position: relative;
    overflow: hidden;
    background: linear-gradient(
        125deg, var(--ta-surface-2), var(--ta-surface)
    );
    border: 1px solid var(--ta-border);
    border-radius: 22px;
    padding: clamp(22px, 4vw, 36px);
    margin: 4px 0 28px;
    box-shadow: var(--ta-shadow);
}

.hero::after {
    content: "";
    position: absolute;
    width: 180px;
    height: 180px;
    right: -70px;
    bottom: -100px;
    border: 1px solid var(--ta-border);
    border-radius: 50%;
    pointer-events: none;
}

.eyebrow {
    color: var(--ta-accent) !important;
    font-size: .72rem;
    font-weight: 800;
    letter-spacing: .16em;
    text-transform: uppercase;
    margin-bottom: 10px;
}

.hero-title {
    color: var(--ta-text) !important;
    font-size: clamp(1.7rem, 3vw, 2.4rem);
    line-height: 1.2;
    font-weight: 800;
    letter-spacing: -.045em;
    margin: 5px 0 12px;
}

.hero .muted {
    max-width: 740px;
    font-size: .96rem;
}

/* ---------- KPI cards ---------- */
.metric-card {
    background: var(--ta-surface);
    border: 1px solid var(--ta-border);
    border-radius: var(--ta-radius);
    padding: 20px;
    min-height: 132px;
    box-shadow: var(--ta-shadow);
    transition: transform .18s ease, border-color .18s ease;
}

.metric-card:hover {
    transform: translateY(-3px);
    border-color: var(--ta-accent);
}

.metric-label {
    color: var(--ta-muted) !important;
    text-transform: uppercase;
    letter-spacing: .09em;
    font-size: .7rem;
    font-weight: 750;
}

.metric-value {
    color: var(--ta-text) !important;
    font-size: clamp(1.65rem, 2.5vw, 2.1rem);
    font-weight: 800;
    line-height: 1.55;
    letter-spacing: -.04em;
    overflow-wrap: anywhere;
}

.metric-foot {
    color: var(--ta-accent) !important;
    font-size: .78rem;
    font-weight: 600;
}

/* ---------- Content panels ---------- */
.panel,
.source-card {
    background: var(--ta-surface);
    color: var(--ta-text);
    border: 1px solid var(--ta-border);
    border-radius: var(--ta-radius);
    padding: 20px;
    margin: 8px 0 18px;
}

.source-card {
    min-height: 100px;
    transition: border-color .18s ease, transform .18s ease;
}

.source-card:hover {
    border-color: var(--ta-accent);
    transform: translateY(-2px);
}

.status-pill {
    display: inline-flex;
    align-items: center;
    padding: 6px 11px;
    border-radius: 999px;
    background: color-mix(
        in srgb, var(--ta-success) 12%, transparent
    );
    border: 1px solid color-mix(
        in srgb, var(--ta-success) 42%, transparent
    );
    color: var(--ta-success) !important;
    font-size: .74rem;
    font-weight: 750;
}

/* ---------- Buttons ---------- */
div.stButton > button,
div.stDownloadButton > button {
    min-height: 43px;
    border-radius: 11px;
    padding: 8px 16px;
    font-weight: 700;
    border: 1px solid var(--ta-border);
    color: var(--ta-text) !important;
    background: var(--ta-surface-2) !important;
    transition: border-color .18s ease, filter .18s ease;
}

div.stButton > button:hover,
div.stDownloadButton > button:hover {
    border-color: var(--ta-accent) !important;
    filter: brightness(1.08);
}

div.stButton > button[kind="primary"],
div.stDownloadButton > button[kind="primary"] {
    background: var(--ta-accent) !important;
    border-color: var(--ta-accent) !important;
    color: var(--ta-accent-text) !important;
}

div.stButton > button[kind="primary"] *,
div.stDownloadButton > button[kind="primary"] * {
    color: var(--ta-accent-text) !important;
}

/* ---------- Inputs and dropdown controls ---------- */
[data-baseweb="input"],
[data-baseweb="textarea"],
[data-baseweb="select"] > div,
[data-testid="stTextInput"] input,
[data-testid="stNumberInput"] input,
[data-testid="stTextArea"] textarea,
[data-testid="stDateInput"] input,
[data-testid="stTimeInput"] input {
    background: var(--ta-surface) !important;
    color: var(--ta-text) !important;
    border-color: var(--ta-border) !important;
    border-radius: 10px !important;
}

input,
textarea,
[data-baseweb="select"] *,
[data-baseweb="input"] *,
[data-baseweb="textarea"] * {
    color: var(--ta-text) !important;
}

input::placeholder,
textarea::placeholder {
    color: var(--ta-muted) !important;
    opacity: .9;
}

/* ---------- Opaque menus and popovers ---------- */
div[data-baseweb="popover"],
div[data-baseweb="popover"] > div,
div[data-baseweb="menu"],
[data-testid="stPopover"],
[data-testid="stPopoverBody"],
[role="listbox"],
[role="menu"] {
    background: var(--ta-surface) !important;
    color: var(--ta-text) !important;
    border-color: var(--ta-border) !important;
    box-shadow: var(--ta-shadow) !important;
}

[role="option"],
[role="menuitem"],
[role="listbox"] *,
[role="menu"] *,
[data-baseweb="popover"] * {
    color: var(--ta-text) !important;
}

[role="option"]:hover,
[role="menuitem"]:hover,
[role="menu"] button:hover {
    background: var(--ta-surface-2) !important;
}

/* ---------- Tabs, expanders and tables ---------- */
[data-testid="stExpander"] {
    border: 1px solid var(--ta-border) !important;
    border-radius: 13px !important;
    overflow: hidden;
}

[data-testid="stExpander"] summary {
    color: var(--ta-text) !important;
    background: var(--ta-surface);
}

[data-testid="stTabs"] button {
    color: var(--ta-text) !important;
}

[data-testid="stDataFrame"],
[data-testid="stTable"] {
    border: 1px solid var(--ta-border);
    border-radius: 12px;
    overflow: hidden;
}

/* ---------- File uploader ---------- */
[data-testid="stFileUploader"] section {
    background: var(--ta-surface) !important;
    border: 1px dashed var(--ta-border) !important;
    border-radius: 14px;
}

[data-testid="stFileUploader"] section * {
    color: var(--ta-text) !important;
}

/* ---------- Alerts and links ---------- */
[data-testid="stAlert"],
[data-testid="stNotification"] {
    border-radius: 12px;
}

a {
    color: var(--ta-accent) !important;
}

/* ---------- Annotated video ---------- */
[data-testid="stVideo"] {
    border: 1px solid var(--ta-border);
    border-radius: 14px;
    overflow: hidden;
    background: #000;
    box-shadow: var(--ta-shadow);
}

video {
    display: block;
    width: 100%;
    height: auto;
    background: #000;
}

/* ---------- Responsive layout ---------- */
@media (max-width: 768px) {
    [data-testid="stMainBlockContainer"] {
        padding: 1rem .8rem 2rem !important;
    }

    .hero {
        border-radius: 16px;
        padding: 22px 18px;
        margin-bottom: 20px;
    }

    .metric-card {
        padding: 15px;
        min-height: 110px;
    }

    .metric-value {
        font-size: 1.6rem;
    }
}

@media (prefers-reduced-motion: reduce) {
    *, *::before, *::after {
        transition-duration: .01ms !important;
        animation-duration: .01ms !important;
    }
}
</style>
""", unsafe_allow_html=True)

# --------------------------- Report helpers ---------------------------

def report_files():
    return sorted(
        OUTPUT.glob("counts_*.json"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )


def video_files():
    return sorted(
        OUTPUT.glob("annotated_*"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )


def read_report(path: Path):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None


def get_counts(data):
    counts = data.get("counts", {})
    up = counts.get("UP", {})
    down = counts.get("DOWN", {})
    classes = sorted(set(up) | set(down))
    classes = [name for name in classes if name != "total_crossings"]
    return up, down, classes


def make_excel(data, recording_start=None):
    """Build a formatted workbook from fields the current JSON actually stores."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    directional = wb.create_sheet("Directional Counts")
    metadata = wb.create_sheet("Metadata")

    navy = "132944"
    cyan = "51D7F5"
    pale = "EAF5FC"
    green = "DDF7EB"
    white = "FFFFFF"
    border_color = "D6E1EC"

    def title_row(ws, title, last_col):
        ws.merge_cells(start_row=1, start_column=1,
                       end_row=1, end_column=last_col)
        cell = ws.cell(1, 1, title)
        cell.font = Font(name="Aptos Display", size=17, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[1].height = 34

    def style_header(ws, row, start, end):
        for col in range(start, end + 1):
            cell = ws.cell(row, col)
            cell.font = Font(bold=True, color=white)
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    title_row(summary, "TRAFFIC ANALYTICS — SUMMARY", 2)
    summary.append([])
    summary.append(["Metric", "Value"])
    style_header(summary, 3, 1, 2)

    counts = data.get("counts", {})
    up, down, classes = get_counts(data)
    total = counts.get(
        "total_crossings",
        sum(int(up.get(c, 0)) + int(down.get(c, 0)) for c in classes),
    )

    summary_rows = [
        ("Total recorded crossings", total),
        ("UP crossings", sum(int(up.get(c, 0)) for c in classes)),
        ("DOWN crossings", sum(int(down.get(c, 0)) for c in classes)),
        ("Video frames processed", data.get("total_frames", "Not available")),
        ("Processing duration (seconds)", data.get("elapsed_seconds", "Not available")),
        ("Report generated at (UTC)", data.get("timestamp_utc", "Not available")),
        ("Recording start time", recording_start.isoformat(sep=" ") if recording_start else "Not supplied"),
    ]
    for row in summary_rows:
        summary.append(list(row))

    for row in range(4, summary.max_row + 1):
        summary.cell(row, 1).fill = PatternFill("solid", fgColor=pale)
        summary.cell(row, 1).font = Font(bold=True, color=navy)

    title_row(directional, "DIRECTIONAL VEHICLE COUNTS", max(3, len(classes) + 1))
    directional.append([])
    directional.append(["Direction"] + classes + ["Total"])
    style_header(directional, 3, 1, len(classes) + 2)

    for direction, values in (("UP", up), ("DOWN", down)):
        row_values = [direction]
        row_values.extend(int(values.get(c, 0)) for c in classes)
        row_values.append(sum(int(values.get(c, 0)) for c in classes))
        directional.append(row_values)

    total_row = ["Combined"]
    total_row.extend(
        int(up.get(c, 0)) + int(down.get(c, 0)) for c in classes
    )
    total_row.append(total)
    directional.append(total_row)

    for cell in directional[directional.max_row]:
        cell.fill = PatternFill("solid", fgColor=green)
        cell.font = Font(bold=True, color=navy)

    title_row(metadata, "ANALYSIS METADATA", 2)
    metadata.append([])
    metadata.append(["Field", "Value"])
    style_header(metadata, 3, 1, 2)

    metadata_rows = [
        ("Source recorded in report", data.get("source", "Not available")),
        ("Report timestamp (UTC)", data.get("timestamp_utc", "Not available")),
        ("Recording start time supplied by user",
         recording_start.isoformat(sep=" ") if recording_start else "Not supplied"),
        ("Total frames", data.get("total_frames", "Not available")),
        ("Elapsed processing seconds", data.get("elapsed_seconds", "Not available")),
        ("Timestamp note",
         "Per-vehicle crossing timestamps are not available in the current JSON report."),
    ]
    for row in metadata_rows:
        metadata.append(list(row))

    for ws in wb.worksheets:
        ws.freeze_panes = "A4"
        ws.sheet_view.showGridLines = False
        for row in ws.iter_rows():
            for cell in row:
                cell.border = Border(
                    bottom=Side(style="hair", color=border_color)
                )
                cell.alignment = Alignment(
                    vertical="center", wrap_text=True
                )
        for col in range(1, ws.max_column + 1):
            letter = get_column_letter(col)
            max_length = max(
                len(str(ws.cell(row, col).value or ""))
                for row in range(1, ws.max_row + 1)
            )
            ws.column_dimensions[letter].width = min(max(max_length + 3, 15), 48)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def safe_rtsp_url(value):
    """Avoid displaying a URL containing credentials in the interface."""
    try:
        parsed = urlsplit(value)
        if parsed.username or parsed.password:
            return f"{parsed.scheme}://[credentials hidden]@{parsed.hostname or ''}"
        return value
    except Exception:
        return "RTSP stream"


def metric(label, value, foot):
    st.markdown(
        f"""
        <div class="metric-card">
            <div class="metric-label">{label}</div>
            <div class="metric-value">{value}</div>
            <div class="metric-foot">{foot}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )


# --------------------------- Sidebar ---------------------------

with st.sidebar:
    st.markdown("## 🚦 Traffic Analytics")
    st.caption("AI traffic intelligence")
    st.markdown("---")

    page = st.radio(
        "NAVIGATION",
        ["Overview", "New Analysis", "Reports History", "Settings"],
        label_visibility="visible",
    )

    st.markdown("---")
    st.markdown("**Pipeline status**")
    st.markdown('<span class="status-pill">● LOCAL ENGINE</span>',
                unsafe_allow_html=True)
    st.caption("Processing runs on the machine hosting this app.")
    st.caption(f"Reports folder: `{OUTPUT}`")


# --------------------------- Overview ---------------------------

if page == "Overview":
    st.markdown(
        """
        <div class="hero">
            <div class="eyebrow">TRAFFIC INTELLIGENCE PLATFORM</div>
            <div class="hero-title">See the flow. Understand the traffic.</div>
            <div class="muted">
                Analyze vehicle movement, inspect directional counts,
                and export structured reports from your videos.
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    reports = report_files()
    latest = read_report(reports[0]) if reports else None

    if latest:
        counts = latest.get("counts", {})
        up, down, classes = get_counts(latest)
        total = counts.get(
            "total_crossings",
            sum(int(up.get(c, 0)) + int(down.get(c, 0)) for c in classes),
        )
        st.caption("LATEST SAVED ANALYSIS")
        c1, c2, c3, c4 = st.columns(4)
        with c1:
            metric("Total crossings", f"{total:,}", "Recorded line crossings")
        with c2:
            metric("UP direction", f"{sum(int(up.get(c, 0)) for c in classes):,}",
                   "Crossings in UP direction")
        with c3:
            metric("DOWN direction", f"{sum(int(down.get(c, 0)) for c in classes):,}",
                   "Crossings in DOWN direction")
        with c4:
            metric("Frames processed", f"{int(latest.get('total_frames', 0)):,}",
                   "Frames reported by pipeline")

        st.markdown("### Vehicle breakdown")
        rows = [
            {
                "Vehicle": name.title(),
                "UP": int(up.get(name, 0)),
                "DOWN": int(down.get(name, 0)),
                "Total": int(up.get(name, 0)) + int(down.get(name, 0)),
            }
            for name in classes
        ]
        if rows:
            import pandas as pd
            df = pd.DataFrame(rows)
            st.dataframe(df, use_container_width=True, hide_index=True)
            chart_df = df.set_index("Vehicle")[["UP", "DOWN"]]
            st.bar_chart(chart_df)
        else:
            st.info("The latest report does not contain class counts.")

        st.caption(f"Report: {reports[0].name}")
        if latest.get("timestamp_utc"):
            st.caption(f"Analysis timestamp (UTC): {latest['timestamp_utc']}")
    else:
        st.markdown(
            """
            <div class="panel">
                <h3>Welcome to your analytics workspace</h3>
                <p class="muted">
                    No reports found yet. Start by uploading a traffic video
                    or connecting to a supported stream.
                </p>
            </div>
            """,
            unsafe_allow_html=True,
        )

    st.markdown("### Quick actions")
    q1, q2 = st.columns(2)
    with q1:
        if st.button("＋ Start new analysis", type="primary", use_container_width=True):
            st.session_state["navigate_to_analysis"] = True
            st.rerun()
    with q2:
        if st.button("View report history", use_container_width=True):
            st.session_state["navigate_to_history"] = True
            st.rerun()

    # Navigation shortcuts from overview.
    if st.session_state.pop("navigate_to_analysis", False):
        st.info("Choose **New Analysis** from the sidebar to start processing.")
    if st.session_state.pop("navigate_to_history", False):
        st.info("Choose **Reports History** from the sidebar to browse saved runs.")


# --------------------------- New analysis ---------------------------

elif page == "New Analysis":
    st.markdown(
        """
        <div class="hero">
            <div class="eyebrow">NEW PROCESSING JOB</div>
            <div class="hero-title">Start an analysis</div>
            <div class="muted">
                Select your source, configure optional recording metadata,
                and run the traffic analytics pipeline.
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    st.markdown("### 1. Choose your source")
    source_mode = st.radio(
        "Input source",
        ["Video file", "Device camera", "CCTV / RTSP"],
        horizontal=True,
    )

    uploaded_file = None
    camera_index = 0
    rtsp_url = ""
    source_ready = False

    if source_mode == "Video file":
        st.markdown(
            '<div class="source-card"><b>🎞️ Video analysis</b>'
            '<p class="muted">Upload an existing traffic recording.</p></div>',
            unsafe_allow_html=True,
        )
        uploaded_file = st.file_uploader(
            "Select a video",
            type=["mp4", "avi", "mov", "mkv", "mpg", "mpeg"],
            help="Use a supported video file. Very large files may take longer to upload.",
        )
        source_ready = uploaded_file is not None

    elif source_mode == "Device camera":
        st.markdown(
            '<div class="source-card"><b>📷 Device camera</b>'
            '<p class="muted">Uses a camera visible to the Python backend.</p></div>',
            unsafe_allow_html=True,
        )
        camera_index = st.number_input(
            "Camera device index",
            min_value=0,
            max_value=10,
            value=0,
            step=1,
            help="Usually 0 for the default camera. WSL may not expose your Windows webcam directly.",
        )
        st.warning(
            "If this app is running in WSL, your Windows webcam may not be "
            "available to OpenCV. Test this mode on the backend machine."
        )
        source_ready = True

    else:
        st.markdown(
            '<div class="source-card"><b>📡 CCTV / RTSP stream</b>'
            '<p class="muted">Connect to a camera stream reachable by this machine.</p></div>',
            unsafe_allow_html=True,
        )
        rtsp_url = st.text_input(
            "RTSP stream URL",
            placeholder="rtsp://camera-address:554/stream",
            help="Do not share URLs containing passwords. Stream access depends on camera configuration and network access.",
            type="password",
        )
        source_ready = bool(rtsp_url.strip())

    st.markdown("### 2. Recording metadata")
    add_recording_time = st.checkbox(
        "I know the recording's actual start date and time",
        value=False,
    )
    recording_start = None

    if add_recording_time:
        dcol, tcol = st.columns(2)
        with dcol:
            recording_date = st.date_input("Recording date", value=date.today())
        with tcol:
            recording_time = st.time_input(
                "Recording start time", value=dt_time(12, 0)
            )
        recording_start = datetime.combine(recording_date, recording_time)
        st.caption(
            "This value is user-supplied metadata. The current pipeline does not "
            "yet calculate each vehicle's exact crossing time."
        )

    with st.expander("Advanced processing settings"):
        st.caption("These settings use the existing project configuration.")
        st.code(str(CONFIG), language="text")
        display_output = st.checkbox(
            "Show processing output in the app",
            value=True,
        )
        st.caption(
            "The current processing command uses the configuration in config.yaml. "
            "Detection settings should be changed there unless explicitly exposed by the pipeline."
        )

    st.markdown("### 3. Run")
    run_clicked = st.button(
        "▶  Start traffic analysis",
        type="primary",
        use_container_width=True,
        disabled=not source_ready,
    )

    if run_clicked:
        source_path = None
        temp_source = None
        mode = "video"

        try:
            start_clock = time.time()

            if source_mode == "Video file":
                suffix = Path(uploaded_file.name).suffix or ".mp4"
                temp_source = OUTPUT / f"input_upload_{int(start_clock)}{suffix}"
                temp_source.write_bytes(uploaded_file.getbuffer())
                source_path = str(temp_source)
                mode = "video"

            elif source_mode == "Device camera":
                source_path = str(int(camera_index))
                mode = "webcam"

            else:
                source_path = rtsp_url.strip()
                mode = "rtsp"

            command = [
                sys.executable,
                str(ROOT / "main.py"),
                "--source", source_path,
                "--mode", mode,
                "--config", str(CONFIG),
                "--output-dir", str(OUTPUT),
                "--no-display",
            ]

            with st.status("Running traffic analytics…", expanded=True) as status:
                st.write("Preparing the input source…")
                st.write("Starting the existing detection pipeline…")
                st.code(
                    " ".join(
                        (f'"{part}"' if " " in part else part)
                        for part in command
                        if part != source_path
                    ) + (
                        ' "<hidden RTSP source>"'
                        if mode == "rtsp"
                        else f' "{source_path}"'
                    ),
                    language="bash",
                )

                # Snapshot old reports so stale files cannot be mistaken for this run.
                before = {p.name for p in report_files()}
                result = subprocess.run(
                    command,
                    cwd=str(ROOT),
                    capture_output=True,
                    text=True,
                    timeout=None,
                )

                duration = time.time() - start_clock
                if result.returncode == 0:
                    after = report_files()
                    fresh = [
                        p for p in after
                        if p.name not in before and p.stat().st_mtime >= start_clock - 1
                    ]
                    if not fresh:
                        fresh = [
                            p for p in after
                            if p.stat().st_mtime >= start_clock - 1
                        ]

                    status.update(
                        label="Analysis completed" if fresh else "Process finished; report not found",
                        state="complete" if fresh else "error",
                        expanded=True,
                    )

                    if fresh:
                        st.session_state["latest_run_report"] = str(fresh[0])
                        st.session_state["latest_run_recording_start"] = (
                            recording_start.isoformat() if recording_start else ""
                        )
                        st.session_state["latest_run_source_mode"] = source_mode
                        st.session_state["latest_run_duration"] = round(duration, 2)
                        st.success("Processing finished. Review your report below.")
                    else:
                        st.warning(
                            "The command returned successfully, but no new JSON report "
                            "was found. Check the pipeline output below."
                        )
                else:
                    status.update(
                        label="Analysis failed",
                        state="error",
                        expanded=True,
                    )
                    st.error(
                        "The processing pipeline returned an error. Review the logs below."
                    )

            if display_output:
                with st.expander("Processing logs", expanded=(result.returncode != 0)):
                    if result.stdout:
                        st.text_area(
                            "Standard output",
                            result.stdout[-12000:],
                            height=220,
                        )
                    if result.stderr:
                        st.text_area(
                            "Error / diagnostic output",
                            result.stderr[-12000:],
                            height=220,
                        )

            if result.returncode != 0:
                st.error(f"Pipeline exit code: {result.returncode}")

        except subprocess.TimeoutExpired:
            st.error("The processing command timed out.")
        except Exception as exc:
            st.error(f"Could not run analysis: {exc}")
        finally:
            # Retain uploaded input files for reproducibility; don't delete automatically.
            pass

    # Show only the report saved by a completed run in this session.
    latest_run_path = st.session_state.get("latest_run_report")
    if latest_run_path and Path(latest_run_path).exists():
        report_path = Path(latest_run_path)
        report_data = read_report(report_path)

        if report_data:
            st.markdown("---")
            st.markdown("## Analysis results")
            counts = report_data.get("counts", {})
            up, down, classes = get_counts(report_data)
            total = counts.get(
                "total_crossings",
                sum(int(up.get(c, 0)) + int(down.get(c, 0)) for c in classes),
            )

            a, b, c = st.columns(3)
            with a:
                metric("Total crossings", f"{total:,}", "Recorded line crossings")
            with b:
                metric("UP", f"{sum(int(up.get(k, 0)) for k in classes):,}",
                       "Directional count")
            with c:
                metric("DOWN", f"{sum(int(down.get(k, 0)) for k in classes):,}",
                       "Directional count")

            import pandas as pd
            df = pd.DataFrame([
                {
                    "Vehicle": k.title(),
                    "UP": int(up.get(k, 0)),
                    "DOWN": int(down.get(k, 0)),
                    "Total": int(up.get(k, 0)) + int(down.get(k, 0)),
                }
                for k in classes
            ])
            if not df.empty:
                st.dataframe(df, use_container_width=True, hide_index=True)
                st.bar_chart(df.set_index("Vehicle")[["UP", "DOWN"]])

            st.caption(f"JSON report: {report_path.name}")

            rec_iso = st.session_state.get("latest_run_recording_start", "")
            rec_start = datetime.fromisoformat(rec_iso) if rec_iso else None
            try:
                excel_bytes = make_excel(report_data, rec_start)
                st.download_button(
                    "⬇ Download Excel report",
                    data=excel_bytes,
                    file_name=f"{report_path.stem}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    type="primary",
                    use_container_width=True,
                )
            except ImportError:
                st.warning(
                    "Excel export needs openpyxl. Install it with: "
                    "`python -m pip install openpyxl`."
                )

            st.download_button(
                "⬇ Download JSON report",
                data=json.dumps(report_data, indent=2).encode("utf-8"),
                file_name=report_path.name,
                mime="application/json",
                use_container_width=True,
            )

            # Associate the newest annotated output created around this run.
            candidates = video_files()
            run_time = report_path.stat().st_mtime
            candidates = [
                p for p in candidates
                if p.stat().st_mtime >= run_time - 600
            ]
            if candidates:
                annotated = candidates[0]
                st.markdown("### Annotated video")
                st.video(str(annotated))
                with open(annotated, "rb") as video_handle:
                    st.download_button(
                        "⬇ Download annotated video",
                        data=video_handle.read(),
                        file_name=annotated.name,
                        mime="video/mp4",
                        use_container_width=True,
                    )


# --------------------------- Report history ---------------------------

elif page == "Reports History":
    st.markdown(
        """
        <div class="hero">
            <div class="eyebrow">SAVED OUTPUTS</div>
            <div class="hero-title">Reports history</div>
            <div class="muted">
                Inspect previous analyses and download their available reports.
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    reports = report_files()
    if not reports:
        st.info("No reports saved yet. Run a video analysis first.")
    else:
        import pandas as pd
        history_rows = []
        for path in reports:
            data = read_report(path)
            if not data:
                continue
            counts = data.get("counts", {})
            up, down, classes = get_counts(data)
            total = counts.get(
                "total_crossings",
                sum(int(up.get(k, 0)) + int(down.get(k, 0)) for k in classes),
            )
            history_rows.append({
                "Report": path.name,
                "Generated (UTC)": data.get("timestamp_utc", ""),
                "Total crossings": total,
                "Frames": data.get("total_frames", ""),
                "Processing seconds": data.get("elapsed_seconds", ""),
            })

        if history_rows:
            st.dataframe(
                pd.DataFrame(history_rows),
                use_container_width=True,
                hide_index=True,
            )

        chosen = st.selectbox(
            "Select a report to inspect",
            options=reports,
            format_func=lambda p: p.name,
        )
        selected_data = read_report(chosen) if chosen else None

        if selected_data:
            st.markdown("### Selected report")
            st.json(selected_data, expanded=False)

            col1, col2 = st.columns(2)
            with col1:
                st.download_button(
                    "Download JSON",
                    data=json.dumps(selected_data, indent=2).encode("utf-8"),
                    file_name=chosen.name,
                    mime="application/json",
                    use_container_width=True,
                )
            with col2:
                try:
                    excel_bytes = make_excel(selected_data)
                    st.download_button(
                        "Download Excel",
                        data=excel_bytes,
                        file_name=f"{chosen.stem}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        use_container_width=True,
                    )
                except ImportError:
                    st.warning("Install openpyxl to enable Excel exports.")

            matching_video = [
                p for p in video_files()
                if p.stat().st_mtime >= chosen.stat().st_mtime - 600
            ]
            if matching_video:
                selected_video = matching_video[0]
                st.markdown("### Annotated video")
                st.video(str(selected_video))
                with open(selected_video, "rb") as handle:
                    st.download_button(
                        "Download annotated video",
                        data=handle.read(),
                        file_name=selected_video.name,
                        mime="video/mp4",
                        use_container_width=True,
                    )


# --------------------------- Settings ---------------------------

elif page == "Settings":
    st.markdown(
        """
        <div class="hero">
            <div class="eyebrow">WORKSPACE CONFIGURATION</div>
            <div class="hero-title">Settings</div>
            <div class="muted">
                Inspect the current local configuration and output locations.
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    st.markdown("### Project")
    st.code(str(ROOT), language="text")
    st.markdown("### Configuration")
    if CONFIG.exists():
        st.code(CONFIG.read_text(encoding="utf-8"), language="yaml")
    else:
        st.warning("config.yaml was not found.")

    st.markdown("### Output directory")
    st.code(str(OUTPUT), language="text")
    st.caption(
        "Output files are stored on the machine running this application. "
        "If running in WSL, open them from Windows using the WSL network path "
        "shown in your terminal or File Explorer."
    )

    st.markdown("### Camera compatibility")
    st.info(
        "A webcam must be accessible to the Python backend. When Streamlit runs "
        "inside WSL, the Windows webcam may not be exposed to OpenCV. CCTV/RTSP "
        "requires a valid stream URL and network access."
    )

st.markdown("---")
st.markdown(
    '<div class="muted" style="text-align:center;padding:8px 0;">'
    'Traffic Analytics · Local processing · Reports saved to output/'
    '</div>',
    unsafe_allow_html=True,
)
